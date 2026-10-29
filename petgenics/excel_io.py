from __future__ import annotations

import io
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

PRODUCT_FIELDS = {
    "product_code",
    "name",
    "brand",
    "family",
    "mrp",
    "pack_size",
    "selling_rate",
    "stock",
}
PARTY_FIELDS = {"party_code", "name"}


def _detect_header_row(raw: pd.DataFrame) -> int:
    limit = min(len(raw), 30)
    candidates = []
    for index in range(limit):
        values = [value for value in raw.iloc[index].tolist() if value is not None and not pd.isna(value)]
        strings = [value for value in values if isinstance(value, str) and value.strip()]
        if strings:
            hints = {
                "code", "product code", "products code", "party code", "name", "product name",
                "party name", "description", "product description", "brand", "category", "sku",
                "item", "pack", "mrp", "qty", "quantity", "stock", "rate", "family", "company",
            }
            normalized = [" ".join(value.casefold().split()) for value in strings]
            has_header_hint = any(
                value in hints
                or value.startswith(("sku (", "pack ", "mrp ", "rate ", "product code ", "product name ", "product description "))
                for value in normalized
            )
            candidates.append((len(strings) * 3 + len(values) + (10 if has_header_hint else 0), -index, index))
    if not candidates:
        raise ValueError("Could not detect a header row. Add column names or specify the header row.")
    return max(candidates)[2]


def _normalize_headers(values: list[Any]) -> list[str]:
    result = []
    counts: dict[str, int] = {}
    for index, value in enumerate(values):
        text = str(value).strip() if value is not None and not pd.isna(value) else ""
        text = text or f"Unnamed: {index}"
        key = text.casefold()
        counts[key] = counts.get(key, 0) + 1
        result.append(text if counts[key] == 1 else f"{text} ({counts[key]})")
    return result


def _read_frame(
    content: bytes,
    filename: str,
    sheet_name: str | int = 0,
    header_row: int | None = None,
) -> pd.DataFrame:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise ValueError("Upload an .xlsx or .xls workbook.")
    try:
        raw = pd.read_excel(io.BytesIO(content), sheet_name=sheet_name, dtype=object, header=None)
        row_index = _detect_header_row(raw) if header_row is None else header_row - 1
        if row_index < 0 or row_index >= len(raw):
            raise ValueError("Header row is outside the worksheet.")
        headers = _normalize_headers(raw.iloc[row_index].tolist())
        frame = raw.iloc[row_index + 1:].copy()
        frame.columns = headers
        frame.index = range(row_index + 2, row_index + 2 + len(frame))
        return frame
    except ImportError as exc:
        raise ValueError("The Excel format reader is not installed.") from exc
    except Exception as exc:
        raise ValueError(f"Could not read the Excel workbook: {exc}") from exc


def preview_workbook(
    content: bytes, filename: str, sheet_name: str | int = 0, header_row: int | None = None
) -> dict[str, Any]:
    frame = _read_frame(content, filename, sheet_name, header_row)
    frame = frame.dropna(axis=0, how="all").dropna(axis=1, how="all")
    columns = [str(column).strip() if not pd.isna(column) else "" for column in frame.columns]
    samples = frame.head(5).where(pd.notna(frame), None).to_dict(orient="records")
    return {
        "columns": columns,
        "header_row": frame.index.min() - 1 if not frame.empty else 1,
        "row_count": len(frame),
        "preview": [{str(k): _json_value(v) for k, v in row.items()} for row in samples],
    }


def _json_value(value: Any) -> Any:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, (pd.Timestamp,)):
        return value.isoformat()
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _read_code_values(
    content: bytes,
    filename: str,
    sheet_name: str | int = 0,
    header_row: int | None = None,
    code_headers: set[str] | None = None,
) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    candidates = code_headers or {
        "product_code", "code", "product code", "products code", "mat_number", "party_code", "party code"
    }
    code_names = {item.casefold() for item in candidates}
    if suffix == ".xlsx":
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook[sheet_name] if isinstance(sheet_name, str) else workbook.worksheets[sheet_name]
        rows = list(sheet.iter_rows(values_only=False))
        if not rows:
            return []
        row_index = _detect_header_row(pd.DataFrame([[cell.value for cell in row] for row in rows])) if header_row is None else header_row - 1
        headers = _normalize_headers([cell.value for cell in rows[row_index]])
        result: list[dict[str, Any]] = []
        for row in rows[row_index + 1:]:
            record: dict[str, Any] = {}
            for header, cell in zip(headers, row):
                value = cell.value
                if header and header.casefold() in code_names and isinstance(value, (int, float)):
                    fmt = cell.number_format or ""
                    if isinstance(value, float) and value.is_integer():
                        value = int(value)
                    zero_mask = re.fullmatch(r"0+", fmt)
                    value = str(value).zfill(len(fmt)) if zero_mask and isinstance(value, int) else str(value)
                record[header] = value
            record["_row"] = len(result) + row_index + 2
            result.append(record)
        return result
    if suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
        sheet = workbook.sheet_by_index(sheet_name) if isinstance(sheet_name, int) else workbook.sheet_by_name(sheet_name)
        raw_rows = [[sheet.cell_value(row, col) for col in range(sheet.ncols)] for row in range(sheet.nrows)]
        row_index = _detect_header_row(pd.DataFrame(raw_rows)) if header_row is None else header_row - 1
        headers = _normalize_headers(raw_rows[row_index])
        result = []
        for row_index_data in range(row_index + 1, sheet.nrows):
            record = {}
            for column, header in enumerate(headers):
                value = sheet.cell_value(row_index_data, column)
                cell = sheet.cell(row_index_data, column)
                if header.casefold() in code_names and cell.ctype == xlrd.XL_CELL_NUMBER:
                    xf = workbook.xf_list[cell.xf_index]
                    fmt = workbook.format_map[xf.format_key].format_str
                    if float(value).is_integer():
                        value = int(value)
                    zero_mask = re.fullmatch(r"0+", fmt)
                    value = str(value).zfill(len(fmt)) if zero_mask and isinstance(value, int) else str(value)
                record[header] = value
            record["_row"] = row_index_data + 1
            result.append(record)
        return result
    return []


def read_mapped_rows(
    content: bytes,
    filename: str,
    mapping: dict[str, str],
    sheet_name: str | int = 0,
    header_row: int | None = None,
) -> list[dict[str, Any]]:
    if not mapping or any(
        target not in PRODUCT_FIELDS | PARTY_FIELDS | {"telegram_user_id", "brand", "display_name"}
        for target in mapping.values()
    ):
        raise ValueError("The column mapping contains unsupported target fields.")
    frame = _read_frame(content, filename, sheet_name, header_row)
    code_headers = {
        source for source, target in mapping.items() if target in {"product_code", "party_code"}
    }
    rows = _read_code_values(content, filename, sheet_name, header_row, code_headers)
    if not rows:
        rows = frame.to_dict(orient="records")
    available = {str(column).strip(): column for column in frame.columns}
    normalized_mapping = {}
    for source, target in mapping.items():
        if source not in available:
            raise ValueError(f"Mapped Excel column does not exist: {source}")
        normalized_mapping[source] = target
    output = []
    for position, raw in enumerate(rows):
        item = {"_row": raw.get("_row", position + 2)}
        for source, target in normalized_mapping.items():
            item[target] = raw.get(source)
        details = {}
        for source, value in raw.items():
            if source == "_row":
                continue
            clean_value = _json_value(value)
            if clean_value not in (None, ""):
                details[source] = clean_value
        item["_source_details"] = details
        output.append(item)
    return output


def workbook_bytes(headers: list[str], rows: list[list[Any]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="244062")
        cell.data_type = "s"
    for row in rows:
        sheet.append([_excel_value(value) for value in row])
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(max(max(len(str(c.value or "")) for c in column) + 2, 12), 45)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _excel_value(value: Any) -> Any:
    return value


def product_template_bytes() -> bytes:
    headers = ["product_code", "name", "brand", "family", "mrp", "pack_size", "selling_rate", "stock"]
    return workbook_bytes(headers, [["", "", "", "", "", "", "", ""]])


def error_report_bytes(errors: list[dict[str, Any]]) -> bytes:
    return workbook_bytes(["row", "product_code", "error"], [
        [item.get("row"), item.get("product_code", ""), item.get("error", "")] for item in errors
    ])
