import asyncio
from datetime import datetime, timezone
from decimal import Decimal
import io
from types import SimpleNamespace

import httpx
from openpyxl import Workbook
from sqlalchemy import create_mock_engine

from petgenics.api import create_app
from petgenics.admin_ui import inline_keyboard, paginate, parse_report_range, period_range, report_day_bounds
from petgenics.bot_menu import (
    ADMIN_COMMANDS,
    ALL_BUTTON_LABELS,
    DEFAULT_COMMANDS,
    MENU_ACTIONS,
    register_bot_menu,
    reply_keyboard,
)
from petgenics.config import Settings, _normalize_database_url, get_settings
from petgenics.excel_io import preview_workbook, read_mapped_rows, workbook_bytes
from petgenics.imports import duplicate_code_rows
from petgenics.marg import order_rows, validate_mapping
from petgenics.money import free_quantity_value, line_total, order_quantity, rate_value
from petgenics.models import (
    Base,
    Order,
    OrderItem,
    Party,
    Product,
    Salesperson,
    SalespersonParty,
)
from petgenics.scanner import extract_document, match_lines


def test_decimal_order_total_and_discount_rounding():
    assert line_total(Decimal("3"), Decimal("10.10"), Decimal("12.5")) == Decimal("26.51")


def test_neon_database_url_is_normalized_for_asyncpg_ssl():
    url = _normalize_database_url(
        "postgresql://user:secret@ep-example.neon.tech/petgenics"
        "?sslmode=require&channel_binding=require"
    )
    assert url.startswith("postgresql+asyncpg://user:secret@ep-example.neon.tech/petgenics?")
    assert "ssl=require" in url
    assert "sslmode=" not in url
    assert "channel_binding=" not in url


def test_render_postgres_url_is_converted_without_changing_database():
    url = _normalize_database_url("postgres://user:secret@db.internal:5432/petgenics")
    assert url == "postgresql+asyncpg://user:secret@db.internal:5432/petgenics"


def test_admin_telegram_ids_are_parsed_and_deduplicated(monkeypatch):
    monkeypatch.setenv("DATABASE_URL", "postgresql://test:test@localhost/petgenics")
    monkeypatch.setenv("ADMIN_TELEGRAM_IDS", "12345,67890,12345")

    assert get_settings().admin_telegram_ids == (12345, 67890)


def test_report_dates_use_ist_boundaries_and_accept_two_formats():
    assert parse_report_range("2026-10-29") == (
        datetime(2026, 10, 29).date(), datetime(2026, 10, 29).date()
    )
    assert parse_report_range("29-10-2026 to 30-10-2026") == (
        datetime(2026, 10, 29).date(), datetime(2026, 10, 30).date()
    )
    assert report_day_bounds(datetime(2026, 10, 30).date(), 330) == (
        datetime(2026, 10, 29, 18, 30, tzinfo=timezone.utc),
        datetime(2026, 10, 30, 18, 30, tzinfo=timezone.utc),
    )


def test_report_range_rejects_reversed_or_oversized_ranges():
    for value in ("2026-10-30 to 2026-10-29", "2024-01-01 to 2026-10-29"):
        try:
            parse_report_range(value)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid report range was accepted.")


def test_report_periods_and_pagination():
    today = datetime(2026, 10, 30).date()
    assert period_range("last7", today) == (
        datetime(2026, 10, 24).date(), today
    )
    assert period_range("lastmonth", today) == (
        datetime(2026, 9, 1).date(), datetime(2026, 9, 30).date()
    )
    page, index, pages = paginate(list(range(19)), 4, 8)
    assert (page, index, pages) == ([16, 17, 18], 2, 3)


def test_telegram_commands_and_reply_keyboard_fit_platform_limits():
    for command in DEFAULT_COMMANDS + ADMIN_COMMANDS:
        assert 1 <= len(command.command) <= 32
        assert all(character.islower() or character.isdigit() or character == "_" for character in command.command)
        assert 3 <= len(command.description) <= 256
    assert len(ALL_BUTTON_LABELS) == len(set(ALL_BUTTON_LABELS))
    assert set(ALL_BUTTON_LABELS) == set(MENU_ACTIONS)
    for keyboard in (reply_keyboard(False), reply_keyboard(True)):
        assert keyboard.resize_keyboard
        assert keyboard.is_persistent


def test_bot_command_menu_registers_user_scopes_and_chat_menu_button():
    class FakeBot:
        def __init__(self):
            self.command_calls = []
            self.menu_button = None

        async def set_my_commands(self, commands, scope):
            self.command_calls.append((commands, scope))

        async def set_chat_menu_button(self, menu_button):
            self.menu_button = menu_button

    bot = FakeBot()
    asyncio.run(register_bot_menu(bot, (123, 456)))
    assert len(bot.command_calls) == 3
    assert bot.command_calls[0][0] == DEFAULT_COMMANDS
    assert bot.command_calls[1][0] == ADMIN_COMMANDS
    assert bot.command_calls[1][1].chat_id == 123
    assert bot.command_calls[2][1].chat_id == 456
    assert bot.menu_button is not None


def test_scanner_admin_callback_passes_context_to_action(monkeypatch):
    from petgenics import admin_reports

    calls = []

    async def fake_scanner_action(query, context, data):
        calls.append((query, context, data))

    monkeypatch.setattr(admin_reports, "scanner_action", fake_scanner_action)
    query = SimpleNamespace()
    context = SimpleNamespace(bot_data={}, user_data={})
    update = SimpleNamespace(callback_query=query)

    assert asyncio.run(
        admin_reports.route_report_callback(update, context, "admin:scanner:purge:confirm")
    )
    assert calls == [(query, context, "admin:scanner:purge:confirm")]


def test_admin_callbacks_reject_unconfigured_telegram_accounts():
    from petgenics.admin_telegram import admin_callback

    class FakeQuery:
        def __init__(self):
            self.answered = False
            self.message = None
            self.edited_text = None

        async def answer(self):
            self.answered = True

        async def edit_message_text(self, text):
            self.edited_text = text

    query = FakeQuery()
    update = SimpleNamespace(
        callback_query=query,
        effective_user=SimpleNamespace(id=999),
    )
    context = SimpleNamespace(bot_data={"admin_telegram_ids": {123}}, user_data={})
    asyncio.run(admin_callback(update, context))
    assert query.answered
    assert query.edited_text == "Administrator access is required."


def test_new_order_asks_for_company_before_parties(monkeypatch):
    from petgenics import telegram_bot

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def commit(self):
            pass

    class FakeMessage:
        def __init__(self):
            self.replies = []

        async def reply_text(self, text, reply_markup=None):
            self.replies.append((text, reply_markup))

    saved = SimpleNamespace(state={})
    salesperson = SimpleNamespace(id=7)
    message = FakeMessage()
    update = SimpleNamespace(
        effective_user=SimpleNamespace(id=123),
        effective_message=message,
    )
    context = SimpleNamespace(bot_data={"admin_telegram_ids": set()}, user_data={})

    async def fake_salesperson(session, user_id):
        return salesperson

    async def fake_saved_state(session, user_id):
        return saved

    async def fake_brands(session, salesperson_id):
        return ["Company A", "Company B"]

    async def fake_persist(session, row, state):
        row.state = state

    monkeypatch.setattr(telegram_bot, "SessionLocal", FakeSession)
    monkeypatch.setattr(telegram_bot, "_salesperson", fake_salesperson)
    monkeypatch.setattr(telegram_bot, "_saved_state", fake_saved_state)
    monkeypatch.setattr(telegram_bot, "accessible_brands", fake_brands)
    monkeypatch.setattr(telegram_bot, "_persist_state", fake_persist)

    asyncio.run(telegram_bot._order(update, context))

    assert saved.state["step"] == "brand_select"
    assert saved.state["brand_choices"] == ["Company A", "Company B"]
    assert message.replies[-1][0] == "First select the company / brand:"
    callbacks = [
        button.callback_data
        for row in message.replies[-1][1].inline_keyboard
        for button in row
    ]
    assert callbacks == ["brand:0", "brand:1"]


def test_company_choice_opens_party_list_with_add_party_action(monkeypatch):
    from petgenics import telegram_bot

    class FakeSession:
        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def commit(self):
            pass

    saved = SimpleNamespace(state={
        "step": "brand_select",
        "brand_choices": ["Company A"],
        "order_start": "company_first",
    })
    party = Party(id=99, party_code="P-1", name="Assigned Shop", active=True)
    query = SimpleNamespace(
        data="brand:0",
        answers=0,
        edits=[],
    )

    async def answer():
        query.answers += 1

    async def edit_message_text(text, reply_markup=None):
        query.edits.append((text, reply_markup))

    query.answer = answer
    query.edit_message_text = edit_message_text
    update = SimpleNamespace(
        callback_query=query,
        effective_user=SimpleNamespace(id=123),
    )
    context = SimpleNamespace(
        bot_data={"admin_telegram_ids": set()},
        user_data={},
    )

    async def fake_salesperson(session, user_id):
        return SimpleNamespace(id=7)

    async def fake_saved_state(session, user_id):
        return saved

    async def fake_brands(session, salesperson_id):
        return ["Company A"]

    async def fake_parties(session, salesperson_id):
        return [party]

    async def fake_persist(session, row, state):
        row.state = state

    monkeypatch.setattr(telegram_bot, "SessionLocal", FakeSession)
    monkeypatch.setattr(telegram_bot, "_salesperson", fake_salesperson)
    monkeypatch.setattr(telegram_bot, "_saved_state", fake_saved_state)
    monkeypatch.setattr(telegram_bot, "accessible_brands", fake_brands)
    monkeypatch.setattr(telegram_bot, "accessible_parties", fake_parties)
    monkeypatch.setattr(telegram_bot, "_persist_state", fake_persist)

    asyncio.run(telegram_bot._callback(update, context))

    assert saved.state["brand"] == "Company A"
    assert saved.state["step"] == "party_search"
    callbacks = [
        button.callback_data
        for row in query.edits[-1][1].inline_keyboard
        for button in row
    ]
    assert callbacks == ["party:99", "party:add"]


def test_new_party_is_active_and_assigned_only_to_requesting_salesperson(monkeypatch):
    from petgenics import telegram_bot

    class FakeSession:
        def __init__(self):
            self.added = []

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return False

        async def scalar(self, statement):
            return None

        def add(self, instance):
            self.added.append(instance)
            if isinstance(instance, Party):
                instance.id = 99

        async def flush(self):
            pass

        async def commit(self):
            pass

    fake_session = FakeSession()
    saved = SimpleNamespace(state={
        "step": "party_add_name",
        "order_start": "company_first",
        "brand": "Company A",
    })
    salesperson = SimpleNamespace(id=7)
    message = SimpleNamespace(text="New Pet Shop", replies=[])

    async def reply_text(text, reply_markup=None):
        message.replies.append(text)

    message.reply_text = reply_text
    update = SimpleNamespace(
        effective_user=SimpleNamespace(id=123),
        effective_message=message,
    )
    context = SimpleNamespace(
        bot_data={"admin_telegram_ids": set()},
        user_data={},
    )

    async def fake_salesperson(session, user_id):
        return salesperson

    async def fake_saved_state(session, user_id):
        return saved

    async def fake_brands(session, salesperson_id):
        return ["Company A"]

    async def fake_persist(session, row, state):
        row.state = state

    async def fake_create_draft(session, user_id, party_id):
        return SimpleNamespace(id=42)

    monkeypatch.setattr(telegram_bot, "SessionLocal", lambda: fake_session)
    monkeypatch.setattr(telegram_bot, "_salesperson", fake_salesperson)
    monkeypatch.setattr(telegram_bot, "_saved_state", fake_saved_state)
    monkeypatch.setattr(telegram_bot, "accessible_brands", fake_brands)
    monkeypatch.setattr(telegram_bot, "_persist_state", fake_persist)
    monkeypatch.setattr(telegram_bot, "create_draft", fake_create_draft)

    asyncio.run(telegram_bot._text(update, context))

    party = next(item for item in fake_session.added if isinstance(item, Party))
    assignment = next(item for item in fake_session.added if isinstance(item, SalespersonParty))
    assert party.name == "New Pet Shop"
    assert party.active is True
    assert party.party_code.startswith("NEW-")
    assert assignment.salesperson_id == salesperson.id
    assert assignment.party_id == party.id
    assert saved.state["party_id"] == party.id
    assert saved.state["brand"] == "Company A"
    assert saved.state["step"] == "search"
    assert "actual ERP party code" in message.replies[-1]


def test_admin_callback_payloads_obey_telegram_size_limit():
    callbacks = [
        "admin:dashboard",
        "admin:orders:filter:submitted",
        "admin:orders:page:submitted:99",
        "admin:orders:bulk:apply",
        "admin:detail:2147483647",
        "admin:order:reject:2147483647",
        "admin:people:party_toggle:2147483647:2147483647",
        "admin:people:party_add_all:apply:2147483647",
        "admin:scanner:delete:confirm:2147483647",
        "admin:report:lastmonth",
        "admin:product:edit:2147483647:selling_rate",
    ]
    for callback in callbacks:
        assert len(callback.encode("utf-8")) <= 64
    try:
        inline_keyboard([[("Too long", "a" * 65)]])
    except ValueError:
        pass
    else:
        raise AssertionError("Oversized callback data was accepted.")


def test_negative_quantity_and_out_of_range_discount_are_rejected():
    for quantity, discount in [(Decimal("-1"), Decimal("0")), (Decimal("1"), Decimal("101"))]:
        try:
            line_total(quantity, Decimal("10"), discount)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid order calculation was accepted.")


def test_order_quantity_rejects_unrepresentable_precision_and_size():
    for quantity in ["0.0001", "1000000000"]:
        try:
            order_quantity(quantity)
        except ValueError:
            pass
        else:
            raise AssertionError("Unrepresentable quantity was accepted.")


def test_excel_header_detection_and_exact_numeric_product_code():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Brand list - current prices"])
    sheet.append(["Product Code", "Product", "MRP", "Pack", "Brand"])
    sheet.append([123, "Sensitive Adult", 99.5, "1 KG", "Example Brand"])
    sheet["A3"].number_format = "000000"
    content = io.BytesIO()
    workbook.save(content)
    data = content.getvalue()

    preview = preview_workbook(data, "products.xlsx")
    assert preview["header_row"] == 2
    assert preview["columns"][:3] == ["Product Code", "Product", "MRP"]
    rows = read_mapped_rows(data, "products.xlsx", {
        "Product Code": "product_code",
        "Product": "name",
        "Brand": "brand",
    })
    assert rows[0]["product_code"] == "000123"
    assert rows[0]["_row"] == 3


def test_stock_only_workbook_with_a_single_named_header_is_detected():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["PRODUCTS CODE", None])
    sheet.append(["SKU-1", 7])
    content = io.BytesIO()
    workbook.save(content)
    rows = read_mapped_rows(content.getvalue(), "stock.xlsx", {
        "PRODUCTS CODE": "product_code",
        "Unnamed: 1": "stock",
    })
    assert rows[0]["product_code"] == "SKU-1"
    assert rows[0]["stock"] == 7


def test_party_import_mapping_reads_selected_worksheet_and_preserves_codes():
    workbook = Workbook()
    workbook.active.title = "Summary"
    sheet = workbook.create_sheet("Party Data")
    sheet.append(["Party Code", "Party Name"])
    sheet.append([123, "Test Pet Shop"])
    sheet["A2"].number_format = "00000"
    content = io.BytesIO()
    workbook.save(content)
    data = content.getvalue()

    rows = read_mapped_rows(
        data,
        "parties.xlsx",
        {"Party Code": "party_code", "Party Name": "name"},
        sheet_name=1,
    )
    assert rows == [{
        "_row": 2,
        "party_code": "00123",
        "name": "Test Pet Shop",
        "_source_details": {"Party Code": "00123", "Party Name": "Test Pet Shop"},
    }]


def test_duplicate_product_codes_are_flagged_without_dropping_source_rows():
    rows = [
        {"_row": 2, "product_code": "SKU-1"},
        {"_row": 3, "product_code": "SKU-2"},
        {"_row": 4, "product_code": "SKU-1"},
    ]
    assert duplicate_code_rows(rows) == {4}


def test_marg_mapping_requires_configured_headers_and_required_fields():
    columns = {
        "Party Key": "party_code",
        "Order Ref": "order_number",
        "Date": "order_date",
        "Item Key": "product_code",
        "Units": "quantity",
    }
    validate_mapping(columns, ["party_code", "product_code", "quantity"])
    try:
        validate_mapping(columns, ["party_code", "rate"])
    except ValueError:
        pass
    else:
        raise AssertionError("An unmapped required field was accepted.")


def test_marg_export_uses_configured_headers_and_order_values():
    order = Order(
        order_number="PG-TEST-1",
        salesperson_id=1,
        party_id=1,
        brand="Test Brand",
        status="Approved",
        created_at=datetime(2026, 1, 2, tzinfo=timezone.utc),
    )
    order.items.append(OrderItem(
        product_id=1,
        product_code="SKU-1",
        product_name="Test product",
        brand="Test Brand",
        pack_size="1 KG",
        quantity=Decimal("2"),
        free_quantity=Decimal("1"),
        rate=Decimal("10.125"),
        discount_percent=Decimal("5"),
        line_total=Decimal("19.24"),
    ))
    content = order_rows(
        order,
        Party(party_code="P-1", name="Test party"),
        Salesperson(display_name="Test seller"),
        {
            "ERP Party": "party_code",
            "ERP Order": "order_number",
            "ERP Date": "order_date",
            "ERP Item": "product_code",
            "ERP Qty": "quantity",
            "ERP Free": "free_quantity",
            "ERP Rate": "rate",
            "ERP Discount": "discount_percent",
        },
        ["party_code", "order_number", "order_date", "product_code", "quantity"],
    )
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(content), data_only=True).active
    assert [cell.value for cell in sheet[1]] == [
        "ERP Party", "ERP Order", "ERP Date", "ERP Item", "ERP Qty", "ERP Free", "ERP Rate", "ERP Discount"
    ]
    assert sheet["A2"].value == "P-1"
    assert sheet["E2"].value == 2
    assert sheet["F2"].value == 1


def test_scanner_does_not_infer_missing_quantity_or_unmatched_product():
    result = match_lines("unknown item 12", [])
    assert result[0]["product_code"] is None
    assert result[0]["quantity"] is None
    assert result[0]["confidence"] == 0
    assert result[0]["warning"]


def test_digital_pdf_text_is_extracted_across_pages():
    import pymupdf

    document = pymupdf.open()
    document.new_page().insert_text((50, 50), "SKU-1 qty 2")
    document.new_page().insert_text((50, 50), "SKU-2 quantity 3")
    result = extract_document(document.tobytes(), "order.pdf")
    assert len(result["pages"]) == 2
    assert "SKU-1 qty 2" in result["pages"][0]["text"]
    assert "SKU-2 quantity 3" in result["pages"][1]["text"]
    assert result["unreadable_pages"] == []


def test_excel_export_is_a_readable_workbook():
    content = workbook_bytes(["product_code"], [["SKU-1"]])
    assert content.startswith(b"PK")


def test_excel_export_treats_formula_like_strings_as_text():
    from openpyxl import load_workbook

    content = workbook_bytes(["name"], [["=1+1"]])
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    assert workbook.active["A2"].data_type == "s"
    assert workbook.active["A2"].value == "=1+1"


def test_free_quantity_is_nonnegative_and_bounded():
    assert free_quantity_value("2.5") == Decimal("2.5")
    for quantity in ["-1", "1000000000", "0.0001"]:
        try:
            free_quantity_value(quantity)
        except ValueError:
            pass
        else:
            raise AssertionError("Invalid free quantity was accepted.")


def test_high_precision_imported_rate_is_preserved_and_product_codes_are_matched_exactly():
    rate = Decimal("118.83983050847459")
    assert rate_value(rate) == rate
    product = Product(id=1, product_code="SKU-123", name="Test food", brand="Test")
    assert match_lines("SKU-123 qty 2", [product])[0]["product_code"] == "SKU-123"
    assert match_lines("SKU-1234 qty 2", [product])[0]["product_code"] is None


def test_api_health_and_admin_routes_enforce_separate_secrets():
    settings = Settings(
        None, "", None, "postgresql+asyncpg://user:pass@localhost/db",
        "admin-test", "health-test", 10000, 24, 15, "INR", None,
    )
    app = create_app(settings)

    async def exercise_routes():
        async with httpx.AsyncClient(
            transport=httpx.ASGITransport(app=app), base_url="http://test"
        ) as client:
            response = await client.get("/healthz")
            assert response.status_code == 200
            assert response.json() == {"status": "ok"}
            assert (await client.get("/health")).status_code == 401
            assert (await client.get("/admin/orders")).status_code == 401

    asyncio.run(exercise_routes())


def test_postgresql_schema_compiles_order_and_stable_code_constraints():
    statements = []
    engine = create_mock_engine(
        "postgresql://",
        lambda statement, *args, **kwargs: statements.append(
            str(statement.compile(dialect=engine.dialect))
        ),
    )
    Base.metadata.create_all(engine)
    sql = "\n".join(statements)
    assert "CREATE TABLE orders" in sql
    assert "CREATE TABLE order_items" in sql
    assert "CREATE UNIQUE INDEX ix_products_product_code" in sql
    assert "ck_order_item_quantity" in sql
    assert "CREATE TABLE scanner_runs" in sql
